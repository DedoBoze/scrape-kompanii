#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
=============================================================================
Скрипта за собирање податоци за компании во Источна Македонија
Цел: Збогатување на B2B Outreach листа за Enterprise ERP & Сметководство SaaS
=============================================================================
Потребни библиотеки:
    pip install requests beautifulsoup4 openpyxl pandas
=============================================================================
"""

import re
import time
import urllib.parse
from datetime import datetime
import requests
from bs4 import BeautifulSoup
import pandas as pd
import openpyxl

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "mk-MK,mk;q=0.9,en-US;q=0.8,en;q=0.7",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
}

CITIES = [
    "stip",
    "kocani",
    "sveti-nikole",
    "vinica",
    "delcevo",
    "berovo",
    "probistip",
    "radovis",
    "strumica"
]

CATEGORIES = [
    # -------------------------------------------------------------
    # 1. Финансии, сметководство и правни услуги (Enterprise Ниша)
    # -------------------------------------------------------------
    ("smetkovodstveni-agencii-i-uslugi", "Сметководствено биро и книговодство", "Enterprise (Биро)"),
    ("revizija", "Ревизорски куќи и финансиски советници", "Enterprise (Биро)"),
    ("konsalting", "Деловен и даночен консалтинг", "Enterprise (Биро)"),
    ("advokati", "Адвокатски канцеларии и правни услуги", "Professional (Фактурирање)"),
    ("notari", "Нотари", "Professional (Фактурирање)"),

    # -------------------------------------------------------------
    # 2. Трговија и дистрибуција (Магацинско работење, WAC & POS)
    # -------------------------------------------------------------
    ("trgovija-na-golemo", "Трговија на големо и дистрибуција", "Professional (WAC & Магацин)"),
    ("trgovija-na-malo", "Трговија на мало и маркети", "Professional (POS & ТКМ)"),
    ("supermarketi", "Супермаркети и синџири маркети", "Professional (POS & ТКМ)"),
    ("avtodelovi-trgovija", "Трговија со автоделови и опрема", "Professional (POS & ТКМ)"),
    ("gradezni-materijali-trgovija", "Трговија со градежни материјали", "Professional (Магацин & Фактури)"),
    ("bela-tehnika", "Бела техника и потрошувачка електроника", "Professional (Магацин & POS)"),
    ("sanitarija-i-vodovod", "Санитарија, водовод и греење", "Professional (Магацин & Фактури)"),
    ("farmacija-apteki", "Аптеки и медицински материјали", "Professional (POS & Залихи)"),
    ("zemjodelski-apteki", "Земјоделски аптеки и ѓубрива", "Professional (Залихи & POS)"),

    # -------------------------------------------------------------
    # 3. Производство и преработувачка индустрија (Материјали, плати & нормативи)
    # -------------------------------------------------------------
    ("tekstil-i-tekstilni-proizvodi-proizvodstvo", "Текстилна индустрија и конфекции", "Professional / Enterprise"),
    ("prehranbeni-proizvodi-proizvodstvo", "Прехранбена индустрија и преработка", "Professional (Залихи & Плати)"),
    ("pekari", "Пекарска индустрија и производство на леб", "Professional (POS & Магацин)"),
    ("vinarii", "Винарии и производство на пијалоци", "Professional (Залихи & Фактури)"),
    ("metalna-industrija", "Метална индустрија, машиноградба и опрема", "Enterprise (Основни средства & WAC)"),
    ("drvo-drvena-industrija", "Дрвна индустрија и преработка на дрво", "Professional (Магацин & Плати)"),
    ("mebel-proizvodstvo-i-trgovija", "Производство и продажба на мебел", "Professional (Магацин & Фактури)"),
    ("plastika-i-proizvodi-od-plastika", "Производство на пластика и амбалажа", "Professional (Залихи & Плати)"),
    ("zemjodelstvo-proizvodstvo", "Земјоделство, овоштарство и сточарство", "Professional (Основни средства & Плати)"),
    ("proizvodstvo", "Општо индустриско производство", "Professional (Залихи & Плати)"),

    # -------------------------------------------------------------
    # 4. Градежништво, недвижности и инженеринг
    # -------------------------------------------------------------
    ("gradeznistvo", "Градежни претпријатија (високо/нискоградба)", "Enterprise (Основни средства & Плати)"),
    ("arhitekti-i-proektiranje", "Архитектура, проектирање и надзор", "Professional (е-Фактура)"),
    ("agencii-za-nedviznosti", "Агенции за недвижности", "Basic / Professional"),

    # -------------------------------------------------------------
    # 5. Транспорт, шпедиција и логистика
    # -------------------------------------------------------------
    ("transport", "Транспортни претпријатија и превоз", "Professional (е-Фактура & Плати)"),
    ("international-transport-forwarding", "Меѓународен транспорт и логистика", "Professional (Девизно & Банка)"),
    ("spediteri-i-spediterski-uslugi", "Шпедитери и царинско посредување", "Professional (е-Фактура)"),
    ("avtoservisi", "Автосервиси и технички прегледи", "Professional (POS & ТКМ)"),

    # -------------------------------------------------------------
    # 6. ИТ, маркетинг и печатарство
    # -------------------------------------------------------------
    ("kompjuteri-softver-razvoj-i-uslugi", "ИТ софтвер, хардвер и сервиси", "Professional (е-Фактура & Банка)"),
    ("marketing-i-reklamni-agencii", "Маркетинг и дигитални агенции", "Professional (е-Фактура)"),
    ("pecatnici", "Печатници, графички дизајн и амбалажа", "Professional (Магацин & Плати)"),

    # -------------------------------------------------------------
    # 7. Угостителство и туризам
    # -------------------------------------------------------------
    ("restorani", "Ресторани и кетеринг", "Professional (POS & ТКМ)"),
    ("hoteli", "Хотели, мотели и туристичко сместување", "Professional (POS & Фактури)"),
    ("turisticki-agencii", "Туристички агенции", "Professional (Фактурирање)"),

    # -------------------------------------------------------------
    # 8. Здравство и приватни ординации
    # -------------------------------------------------------------
    ("polikliniki", "Приватни здравствени установи и поликлиники", "Professional (е-Фактура & Плати)"),
    ("stomatoloski-ordinacii", "Стоматолошки ординации", "Basic / Professional")
]

EMAIL_REGEX = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_REGEX = re.compile(r"(\+?389\s?\d{1,2}\s?\d{3}\s?\d{3}|\b0\d{1,2}[\s/-]?\d{3}[\s/-]?\d{3}\b)")

def clean_text(text):
    if not text:
        return ""
    return re.sub(r"\s+", " ", text).strip()

def extract_email_from_text(text):
    if not text:
        return ""
    match = EMAIL_REGEX.search(text)
    return match.group(0) if match else ""

def get_company_details(detail_url, session):
    """
    Ја посетува поединечната страница на фирмата за да извлече емаил адреса и веб-страница.
    """
    email = ""
    website = ""
    try:
        resp = session.get(detail_url, headers=HEADERS, timeout=10)
        if resp.status_code == 200:
            soup = BeautifulSoup(resp.text, "html.parser")
            
            # Барање mailto линкови
            mailto = soup.select_one('a[href^="mailto:"]')
            if mailto:
                email = mailto.get("href", "").replace("mailto:", "").split("?")[0].strip()
            
            # Ако нема директен mailto, пребарај текст
            if not email:
                email = extract_email_from_text(resp.text)
            
            # Барање надворешна веб страница на фирмата
            for a_tag in soup.select('a[href^="http"]'):
                href = a_tag.get("href", "")
                if "zk.mk" not in href and "facebook.com" not in href and "google.com" not in href:
                    website = href
                    break
    except Exception as e:
        print(f"    [Грешка при детална страница {detail_url}]: {e}")
    
    return email, website

def scrape_zk_category_city(category_slug, category_name, icp_tier, city, session):
    results = []
    skip = 0
    step = 10
    max_pages = 10  # До 10 страници (100 резултати) по дејност/град

    print(f"\n--- Пребарување: Дејност='{category_name}' во Град='{city.upper()}' ---")

    while skip < (max_pages * step):
        url = f"https://zk.mk/{category_slug}/{city}?skip={skip}"
        try:
            time.sleep(1.2)  # Пауза за заштита од блокирање
            resp = session.get(url, headers=HEADERS, timeout=12)
            if resp.status_code != 200:
                print(f"  [Статус {resp.status_code}] за {url}")
                break

            soup = BeautifulSoup(resp.text, "html.parser")
            
            # Селектирање картички на фирми
            cards = soup.select(".short_details, .company_item, div[itemtype*='LocalBusiness'], .result_item")
            
            if not cards:
                cards = [a.parent for a in soup.select("h3 a, h2 a") if a.get("href", "").startswith("/")]

            if not cards:
                print(f"  Нема повеќе резултати на skip={skip}.")
                break

            found_on_page = 0
            for card in cards:
                title_elem = card.select_one("h3 a, h2 a, a.company_title") or card.find("a")
                if not title_elem or not title_elem.text.strip():
                    continue

                name = clean_text(title_elem.text)
                detail_link = title_elem.get("href", "")
                if detail_link.startswith("/"):
                    detail_link = urllib.parse.urljoin("https://zk.mk", detail_link)

                card_text = card.get_text(separator=" ")
                
                # Екстракција на телефон
                phone_match = PHONE_REGEX.search(card_text)
                phone = phone_match.group(0) if phone_match else ""
                
                # Екстракција на адреса
                addr_match = re.search(r"Адреса\s*:\s*([^·\n\r]+)", card_text, re.IGNORECASE)
                address = clean_text(addr_match.group(1)) if addr_match else ""

                # Емаил од почетната картичка
                email = extract_email_from_text(card_text)
                website = ""

                # Ако нема емаил, отвори ја деталната страница на фирмата
                if not email and detail_link and "zk.mk" in detail_link:
                    time.sleep(0.5)
                    email, website = get_company_details(detail_link, session)

                results.append({
                    "Име на Компанија": name,
                    "Град / Општина": city.capitalize(),
                    "Дејност / Категорија": category_name,
                    "Тип на Клиент (ICP)": "Сметководствено биро" if "сметковод" in category_name.lower() else "Трговија/Производство",
                    "Емаил Адреса": email,
                    "Телефон": phone,
                    "Адреса": address,
                    "Веб-страница": website if website else detail_link,
                    "Препорачан SaaS План": icp_tier,
                    "Статус на Контакт": "Нов",
                    "Датум на Екстракција": datetime.now().strftime("%Y-%m-%d")
                })
                found_on_page += 1

            print(f"  Страница со skip={skip}: Најдени {found_on_page} фирми.")
            if found_on_page == 0:
                break

            skip += step

        except Exception as e:
            print(f"  [Исклучок]: {e}")
            break

    return results

def main():
    session = requests.Session()
    all_companies = []

    print("=" * 70)
    print("Започнува собирањето на податоци за компании во Источна Македонија...")
    print("=" * 70)

    for cat_slug, cat_name, icp_tier in CATEGORIES:
        for city in CITIES:
            data = scrape_zk_category_city(cat_slug, cat_name, icp_tier, city, session)
            all_companies.extend(data)
            print(f"Вкупно собрани досега: {len(all_companies)} компании.")

    if not all_companies:
        print("Не беа пронајдени компании. Проверете ја интернет конекцијата.")
        return

    # Отстранување на дупликати по име и град
    df = pd.DataFrame(all_companies)
    df.drop_duplicates(subset=["Име на Компанија", "Град / Општина"], inplace=True)
    df.insert(0, "Р.Бр.", range(1, len(df) + 1))

    output_filename = "Kompanii_Istocna_Makedonija_Scraped.xlsx"
    
    # Снимање во форматиран Excel
    with pd.ExcelWriter(output_filename, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Собрани_Компании", index=False)
        
        wb = writer.book
        ws = writer.sheets["Собрани_Компании"]
        ws.views.sheetView[0].showGridLines = True
        
        from openpyxl.styles import Font, PatternFill, Alignment
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        
        for col_num in range(1, len(df.columns) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
            col_letter = openpyxl.utils.get_column_letter(col_num)
            max_len = max(df.iloc[:, col_num-1].astype(str).map(len).max(), len(str(cell.value))) + 3
            ws.column_dimensions[col_letter].width = min(max_len, 45)

    print("\n" + "=" * 70)
    print(f"УСПЕХ! Вкупно {len(df)} уникатни компании се зачувани во: {output_filename}")
    print("=" * 70)

if __name__ == "__main__":
    main()